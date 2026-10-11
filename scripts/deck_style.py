"""Shared visual style and slide builders for TruckIDMe decks.

Palette: charcoal + amber (dashboard warning lights), off-white cards.
Used by build_proposal_deck.py and build_lit_deck.py.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

CHARCOAL = RGBColor(0x2B, 0x3A, 0x42)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
AMBER_DK = RGBColor(0xA8, 0x6A, 0x12)
TINT_GREY = RGBColor(0xF2, 0xF3, 0xF4)
TINT_AMBER = RGBColor(0xFD, 0xF4, 0xE6)
TINT_TEAL = RGBColor(0xEC, 0xF4, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x7A, 0x82)
TEAL = RGBColor(0x2E, 0x7D, 0x7B)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def new_deck() -> Presentation:
    """A 16:9 presentation with no inherited placeholders."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def blank(prs: Presentation):
    """Add a slide using the blank layout."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill(slide, color: RGBColor) -> None:
    """Paint the whole slide background."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    """Add an empty text box with no internal padding."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def write(tf, text, *, size, color, bold=False, font=BODY_FONT, space_after=0,
          align=PP_ALIGN.LEFT, first=False, line=None, italic=False):
    """Append a paragraph, honouring **bold** markers inside the text."""
    para = tf.paragraphs[0] if first else tf.add_paragraph()
    para.alignment = align
    para.space_after = Pt(space_after)
    if line:
        para.line_spacing = line
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        run = para.add_run()
        run.text = chunk
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = font
        run.font.bold = bold or (i % 2 == 1)
        run.font.italic = italic
    return para


def title_block(slide, title, kicker=None, *, dark=False, top=Inches(0.6)):
    """Slide title plus optional kicker line."""
    tf = textbox(slide, Inches(0.7), top, Inches(12), Inches(0.75))
    write(tf, title, size=32, color=AMBER if dark else CHARCOAL, bold=True,
          font=HEAD_FONT, first=True)
    if kicker:
        tf2 = textbox(slide, Inches(0.7), top + Inches(0.72), Inches(12), Inches(0.4))
        write(tf2, kicker, size=12.5, color=TINT_GREY if dark else MUTED,
              first=True, italic=True)


def card(slide, x, y, w, h, label, bullets, *, tint=TINT_GREY,
         label_color=CHARCOAL, size=12.5):
    """A titled card holding bullet lines."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = tint
    box.line.fill.background()
    box.shadow.inherit = False
    box.adjustments[0] = 0.04
    box.text_frame.text = ""

    pad = Inches(0.3)
    tf = textbox(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad)
    if label:
        write(tf, label, size=13, color=label_color, bold=True, font=HEAD_FONT,
              space_after=9, first=True)
    for i, b in enumerate(bullets):
        write(tf, "▪  " + b, size=size, color=CHARCOAL, space_after=7, line=1.06,
              first=(i == 0 and not label))


def badge(slide, text, x, y, *, size=Inches(0.52)):
    """Numbered circle motif."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    c.fill.solid()
    c.fill.fore_color.rgb = AMBER
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(tf, str(text), size=17, color=CHARCOAL, bold=True, font=HEAD_FONT,
          align=PP_ALIGN.CENTER, first=True)


def two_card_slide(prs, title, kicker, left, right, *, notes="", n=None):
    """Two side-by-side cards (label, bullets) beneath a title."""
    s = blank(prs)
    fill(s, WHITE)
    if n is not None:
        badge(s, n, Inches(0.55), Inches(0.5))
        tf = textbox(s, Inches(1.25), Inches(0.48), Inches(11.5), Inches(0.62))
        write(tf, title, size=30, color=CHARCOAL, bold=True, font=HEAD_FONT, first=True)
        tf2 = textbox(s, Inches(1.25), Inches(1.08), Inches(11.5), Inches(0.3))
        write(tf2, kicker, size=11.5, color=MUTED, first=True, italic=True)
    else:
        title_block(s, title, kicker)
    top, h = Inches(1.65), Inches(5.3)
    card(s, Inches(0.55), top, Inches(5.95), h, left[0], left[1])
    card(s, Inches(6.83), top, Inches(5.95), h, right[0], right[1],
         tint=TINT_AMBER, label_color=AMBER_DK)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def three_card_slide(prs, title, kicker, cards, *, notes=""):
    """Three cards in a row; each card is (label, [bullets])."""
    s = blank(prs)
    fill(s, WHITE)
    title_block(s, title, kicker)
    tints = [TINT_GREY, TINT_AMBER, TINT_TEAL]
    labels = [CHARCOAL, AMBER_DK, TEAL]
    x, w, gap = Inches(0.55), Inches(4.0), Inches(0.14)
    for i, (label, bullets) in enumerate(cards):
        card(s, x + i * (w + gap), Inches(1.7), w, Inches(5.25), label, bullets,
             tint=tints[i % 3], label_color=labels[i % 3], size=11.5)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def bullet_slide(prs, title, kicker, bullets, *, notes="", dark=False, size=15):
    """Single-column slide."""
    s = blank(prs)
    fill(s, CHARCOAL if dark else WHITE)
    title_block(s, title, kicker, dark=dark, top=Inches(0.75))
    tf = textbox(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(4.8))
    for i, b in enumerate(bullets):
        write(tf, "▪  " + b, size=size, color=WHITE if dark else CHARCOAL,
              space_after=13, line=1.1, first=(i == 0))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def table_slide(prs, title, kicker, headers, rows, *, notes="", col_w=None,
                top=Inches(1.85), row_h=Inches(0.42), body_size=10.5):
    """Slide holding a single table."""
    s = blank(prs)
    fill(s, WHITE)
    title_block(s, title, kicker)
    shape = s.shapes.add_table(len(rows) + 1, len(headers), Inches(0.7), top,
                               Inches(11.9), row_h * (len(rows) + 1))
    table = shape.table
    if col_w:
        total = sum(col_w)
        for i, frac in enumerate(col_w):
            table.columns[i].width = Emu(int(Inches(11.9) * frac / total))
    for c, head in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CHARCOAL
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.text_frame.word_wrap = True
        write(cell.text_frame, head, size=11.5, color=WHITE, bold=True,
              font=HEAD_FONT, first=True)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (TINT_AMBER if r == 1 else
                                        (TINT_GREY if r % 2 == 0 else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.text_frame.word_wrap = True
            write(cell.text_frame, val, size=body_size, color=CHARCOAL, first=True)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s
